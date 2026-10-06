#!/usr/bin/env python3
"""Extracts a monthly-costs workbook (one sheet per month, named like "Octubre 22") into JSON for scripts/import-excel.mjs.

  python3 scripts/extract-excel.py workbook.xlsx out.json --admin "Name" [--skip "label prefix" ...]
  (needs: pip install openpyxl)

  --admin  the row name of the person who pays the bills (not a roommate, so no payments are generated for them)
  --skip   ignore left-table rows whose name starts with this text (notes such as "Total" or "Devolver" are skipped already)

Right-hand table ("Costos"): category, amount, receipt label + Drive hyperlink, debit date, bank.
Left-hand table: one row per person; the roommate's "Total" column is what they paid that month.
"""
import argparse, json
import openpyxl

MONTHS = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6, 'julio': 7, 'agosto': 8,
          'septiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12}
NOT_PEOPLE = {'total', 'devolver'}

def category(label):
    l = label.lower()
    for key, name in (('alquiler', 'Alquiler'), ('luz', 'Luz'), ('internet', 'Internet'), ('tel', 'Teléfono'),
                      ('gas', 'Gas'), ('limpieza', 'Limpieza'), ('agua', 'Agua')):
        if key in l:
            return name
    raise ValueError(label)

def main(src, out, admin, skip):
    wb = openpyxl.load_workbook(src)
    wv = openpyxl.load_workbook(src, data_only=True)
    result = []
    for ws in wb:
        v = wv[ws.title]
        name, yy = ws.title.split()
        month = f'20{yy}-{MONTHS[name.lower()]:02d}'
        hc = next((c for row in ws.iter_rows(min_row=1, max_row=3) for c in row if c.value == 'Costos'), None)
        r0, c0 = hc.row, hc.column
        bills, r = [], r0 + 1
        while ws.cell(r, c0).value not in (None, 'Total'):
            label = ws.cell(r, c0).value
            cell = ws.cell(r, c0 + 2)
            amt = v.cell(r, c0 + 1).value
            date = v.cell(r, c0 + 3).value
            bills.append({
                'label': label, 'category': category(label), 'amount': round(float(amt), 2) if amt else 0,
                'fileLabel': cell.value, 'link': cell.hyperlink.target if cell.hyperlink else None,
                'debitDate': date.strftime('%Y-%m-%d') if date else None, 'bank': v.cell(r, c0 + 4).value,
                'row': r,
            })
            r += 1
        cols = {ws.cell(2, c).value: c for c in range(2, c0) if ws.cell(2, c).value}
        rent_col, total_col = cols['Alquiler'], cols['Total']
        people = []
        for rr in range(3, 40):
            n = ws.cell(rr, 1).value
            if not isinstance(n, str) or n.strip().lower() in NOT_PEOPLE | {admin.lower()} or n.lower().startswith(skip):
                continue
            total, rent = v.cell(rr, total_col).value, v.cell(rr, rent_col).value
            if isinstance(total, (int, float)):
                people.append({'name': n.strip(), 'paid': round(float(total), 2),
                               'rent': round(float(rent), 2) if isinstance(rent, (int, float)) else None})
        result.append({'sheet': ws.title, 'month': month, 'bills': bills, 'people': people})
    json.dump(result, open(out, 'w'), ensure_ascii=False, indent=1)
    print(f'{len(result)} sheets, {sum(len(s["bills"]) for s in result)} bills, {sum(len(s["people"]) for s in result)} payments')

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('workbook')
    ap.add_argument('out')
    ap.add_argument('--admin', required=True)
    ap.add_argument('--skip', action='append', default=[])
    a = ap.parse_args()
    main(a.workbook, a.out, a.admin.strip(), tuple(x.lower() for x in a.skip) or ('\0',))
